"""Bounded local parsing from BytesIO; no files, provider, OCR or semantic inference."""

from contextvars import ContextVar
from io import BytesIO
import logging
import posixpath
import sys
from time import perf_counter
from urllib.parse import unquote
from zipfile import ZipFile, ZIP_STORED, ZIP_DEFLATED

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
from pypdf import PdfReader, apply_configuration

from resume_intake.models import (
    ParsedResumeBlock, ParsedResumeDocument, ResumeParseResult, ResumeParseStatus, UploadedResume,
)
from resume_intake.policy import (
    MAX_UPLOAD_BYTES, MAX_PDF_PAGES, MAX_ARCHIVE_ENTRIES, MAX_EXPANDED_BYTES,
    MAX_PART_BYTES, MAX_COMPRESSION_RATIO, MAX_BLOCKS, MAX_TEXT_CHARACTERS, MIN_TEXT_CHARACTERS,
)


class _Rejected(Exception):
    def __init__(self, code: str, status: ResumeParseStatus = ResumeParseStatus.PARSE_FAILED):
        self.code, self.status = code, status


_PARSING = ContextVar("orange_local_resume_parse", default=False)


class _PrivateParserLogs(logging.Filter):
    def filter(self, record: logging.LogRecord) -> bool:
        return not _PARSING.get()


_LOG_FILTER = _PrivateParserLogs()
# pypdf warnings can contain document objects. Suppress only this parse context,
# not other threads or the application's logging. No raw exception is logged.
for _name in tuple(sys.modules):
    if _name == "pypdf" or _name.startswith("pypdf."):
        logging.getLogger(_name).addFilter(_LOG_FILTER)


def _docx_container(data: bytes) -> None:
    if not data.startswith(b"PK\x03\x04"):
        raise _Rejected("INVALID_SIGNATURE")
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ARCHIVE_ENTRIES:
            raise _Rejected("ARCHIVE_LIMIT")
        names: set[str] = set()
        total = 0
        for entry in entries:
            name = entry.filename
            if (name in names or "\\" in name or name.startswith("/") or
                    any(part in ("", ".", "..") for part in name.rstrip("/").split("/")) or
                    ":" in name or (entry.external_attr >> 16) & 0o170000 == 0o120000):
                raise _Rejected("UNSAFE_ARCHIVE_PATH")
            names.add(name)
            total += entry.file_size
            if (entry.file_size > MAX_PART_BYTES or total > MAX_EXPANDED_BYTES or
                    entry.file_size > max(entry.compress_size, 1) * MAX_COMPRESSION_RATIO):
                raise _Rejected("ARCHIVE_LIMIT")
            if entry.flag_bits & 1 or entry.compress_type not in (ZIP_STORED, ZIP_DEFLATED):
                raise _Rejected("UNSUPPORTED_ARCHIVE")
            lowered = name.casefold()
            if "vba" in lowered or lowered.startswith("word/embeddings/"):
                raise _Rejected("ACTIVE_CONTENT_UNSUPPORTED")
        if not {"[Content_Types].xml", "_rels/.rels", "word/document.xml"} <= names:
            raise _Rejected("INVALID_DOCX_CONTAINER")
        main_type = False
        main_relation = False
        for entry in entries:
            if not entry.filename.endswith((".xml", ".rels")):
                continue
            # ZipFile verifies CRC; read only after all declared expansion bounds.
            raw = archive.read(entry)
            root = etree.fromstring(raw, parser=etree.XMLParser(
                resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False,
            ))
            if root.getroottree().docinfo.doctype or any(isinstance(n, etree._Entity) for n in root.iter()):
                raise _Rejected("ACTIVE_CONTENT_UNSUPPORTED")
            if entry.filename == "[Content_Types].xml":
                for node in root:
                    content_type = node.get("ContentType", "")
                    if any(kind in content_type.casefold() for kind in ("macro", "vba", "oleobject")):
                        raise _Rejected("ACTIVE_CONTENT_UNSUPPORTED")
                    if node.get("PartName") == "/word/document.xml":
                        main_type = content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
            if entry.filename.endswith(".rels"):
                base = entry.filename.rpartition("/_rels/")[0] if "/_rels/" in entry.filename else ""
                for node in root:
                    if node.get("Type", "").rsplit("/", 1)[-1] in {"oleObject", "package", "aFChunk"}:
                        raise _Rejected("ACTIVE_CONTENT_UNSUPPORTED")
                    if node.get("TargetMode") == "External":
                        continue  # Never fetch external hyperlinks/resources.
                    target = unquote(node.get("Target", ""))
                    resolved = posixpath.normpath(posixpath.join(base, target))
                    if "\\" in target or ":" in target or target.startswith("/") or resolved.startswith("../"):
                        raise _Rejected("UNSAFE_ARCHIVE_PATH")
                    if entry.filename == "_rels/.rels" and node.get("Type", "").endswith("/officeDocument"):
                        main_relation = resolved == "word/document.xml"
            if entry.filename == "word/document.xml" and root.tag != "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}document":
                raise _Rejected("INVALID_DOCX_CONTAINER")
        if not main_type or not main_relation:
            raise _Rejected("INVALID_DOCX_CONTAINER")


def _blocks(upload: UploadedResume):
    blocks: list[ParsedResumeBlock] = []
    text_count = 0

    def append(text: str, location: str, kind: str, page: int | None = None) -> None:
        nonlocal text_count
        text = text.strip()
        if not text:
            return
        text_count += len(text)
        if text_count > MAX_TEXT_CHARACTERS or len(blocks) >= MAX_BLOCKS:
            raise _Rejected("TEXT_LIMIT")
        blocks.append(ParsedResumeBlock(
            f"{upload.source_id}:block:{len(blocks) + 1}", upload.source_id,
            upload.file_type, location, kind, text, page,
        ))

    return blocks, append


def _parse_pdf(upload: UploadedResume) -> ParsedResumeDocument:
    if not upload.data.startswith(b"%PDF-"):
        raise _Rejected("INVALID_SIGNATURE")
    with apply_configuration(
        maximum_declared_stream_length=MAX_PART_BYTES,
        zlib_maximum_output_length=MAX_PART_BYTES,
        zlib_maximum_recovery_input_length=MAX_PART_BYTES,
        array_based_stream_maximum_output_length=MAX_PART_BYTES,
        lzw_maximum_output_length=MAX_PART_BYTES,
        run_length_maximum_output_length=MAX_PART_BYTES,
        jbig2_maximum_output_length=MAX_PART_BYTES, jbig2dec_binary=None,
        page_tree_maximum_entries=256, page_tree_maximum_depth=32,
        xform_maximum_invocations_per_extraction=256,
    ):
        reader = PdfReader(BytesIO(upload.data), strict=True)
        if reader.is_encrypted:
            raise _Rejected("ENCRYPTED_PDF", ResumeParseStatus.ENCRYPTED)
        if sum(len(objects) for objects in reader.xref.values()) > 10_000:
            raise _Rejected("OBJECT_LIMIT")
        pages = len(reader.pages)
        if not 0 < pages <= MAX_PDF_PAGES:
            raise _Rejected("PAGE_LIMIT")
        blocks, append = _blocks(upload)
        total = 0
        for number, page in enumerate(reader.pages, 1):
            content = page.get_contents()
            length = len(content.get_data()) if content is not None else 0
            total += length
            if length > MAX_PART_BYTES or total > MAX_EXPANDED_BYTES:
                raise _Rejected("STREAM_LIMIT")
            append(page.extract_text(), f"page/{number}", "page_text", number)
        return ParsedResumeDocument(upload.source_id, "pdf", tuple(blocks), pages)


def _parse_docx(upload: UploadedResume) -> ParsedResumeDocument:
    _docx_container(upload.data)
    document = Document(BytesIO(upload.data))
    blocks, append = _blocks(upload)

    def visit(container, location: str, depth: int = 0) -> None:
        if depth > 12:
            raise _Rejected("DOCUMENT_DEPTH_LIMIT")
        for position, item in enumerate(container.iter_inner_content(), 1):
            here = f"{location}/{position}"
            if isinstance(item, Paragraph):
                append(item.text, here, "paragraph" if depth == 0 else "table_cell")
            elif isinstance(item, Table):
                seen = set()
                for row_position, row in enumerate(item.rows, 1):
                    for cell_position, cell in enumerate(row.cells, 1):
                        if cell._tc in seen:
                            continue
                        seen.add(cell._tc)
                        visit(cell, f"{here}/table/row/{row_position}/cell/{cell_position}", depth + 1)

    visit(document, "body")
    # Header/footer text is optional and remains explicitly located, not body order.
    seen_parts = set()
    for index, section in enumerate(document.sections, 1):
        for label in ("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer"):
            part = getattr(section, label)
            if part.is_linked_to_previous:
                continue
            if part.part.partname in seen_parts:
                continue
            seen_parts.add(part.part.partname)
            visit(part, f"section/{index}/{label}")
    return ParsedResumeDocument(upload.source_id, "docx", tuple(blocks))


def parse_resume(upload: UploadedResume) -> ResumeParseResult:
    """Return only safe codes on failure; discard partial extraction, never repair."""
    started = perf_counter()
    token = _PARSING.set(True)
    try:
        if upload.file_type not in ("pdf", "docx"):
            raise _Rejected("UNSUPPORTED_TYPE", ResumeParseStatus.UNSUPPORTED)
        if not upload.data:
            raise _Rejected("EMPTY_FILE")
        if len(upload.data) > MAX_UPLOAD_BYTES:
            raise _Rejected("UPLOAD_LIMIT", ResumeParseStatus.TOO_LARGE)
        if upload.size_bytes != len(upload.data):
            raise _Rejected("INVALID_SIZE")
        document = _parse_pdf(upload) if upload.file_type == "pdf" else _parse_docx(upload)
        if sum(sum(char.isalnum() for char in block.text) for block in document.blocks) < MIN_TEXT_CHARACTERS:
            raise _Rejected("NO_RELIABLE_TEXT", ResumeParseStatus.SCAN_OR_IMAGE_ONLY)
        return ResumeParseResult(ResumeParseStatus.READY, document=document,
                                 elapsed_ms=(perf_counter() - started) * 1000)
    except _Rejected as error:
        return ResumeParseResult(error.status, error.code, elapsed_ms=(perf_counter() - started) * 1000)
    except Exception:
        return ResumeParseResult(ResumeParseStatus.PARSE_FAILED, "PARSE_FAILED",
                                 elapsed_ms=(perf_counter() - started) * 1000)
    finally:
        _PARSING.reset(token)
