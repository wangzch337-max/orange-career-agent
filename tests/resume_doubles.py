"""Public synthetic documents generated in memory, not personal resume fixtures."""

from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

SYNTHETIC_TEXT = "Public synthetic work experience: maintained schedules and documented quality checks."


def pdf_bytes(*texts: str, encrypted: bool = False) -> bytes:
    writer = PdfWriter()
    for text in texts or (SYNTHETIC_TEXT,):
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 40 700 Td (" + text.encode("ascii") + b") Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("synthetic-only-password")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph(SYNTHETIC_TEXT)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Work responsibilities"
    table.cell(0, 1).text = "Coordinate public synthetic delivery records"
    document.add_paragraph("Professional qualification: synthetic safety course")
    document.sections[0].header.paragraphs[0].text = "Public synthetic document"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def replace_zip(data: bytes, changes: dict[str, bytes]) -> bytes:
    output = BytesIO()
    with ZipFile(BytesIO(data)) as original, ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name in original.namelist():
            archive.writestr(name, changes.pop(name, original.read(name)))
        for name, value in changes.items():
            archive.writestr(name, value)
    return output.getvalue()
