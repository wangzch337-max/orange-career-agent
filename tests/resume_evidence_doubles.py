"""Small public synthetic shape fixtures, not a cross-background evaluator."""

from io import BytesIO
from uuid import uuid4

from docx import Document

from resume_intake.models import ParsedResumeDocument, ParsedResumeBlock

BACKGROUNDS = {
    "student": ("Education: Diploma in Hospitality", "Public Synthetic College", "2024 - 2026", "Coordinated a student welcome event", "English and Cantonese"),
    "audit": ("Audit Associate", "Public Synthetic Audit Firm", "2021 - 2024", "Prepared statutory audit working papers and reconciliation", "Excel and financial reporting"),
    "mechanical": ("Mechanical Engineer", "Public Synthetic Manufacturing", "2018 - 2024", "Reviewed tolerances and coordinated preventive maintenance", "CAD and quality checks"),
    "marketing": ("E-commerce Specialist", "Public Synthetic Retail", "2020 - 2024", "Increased conversion by 18%", "Campaign analytics and customer collaboration"),
    "design": ("UX Designer", "Public Synthetic Design Studio", "2019 - 2024", "Conducted usability interviews and documented accessibility findings", "Figma and prototyping"),
    "switcher": ("Operations Supervisor", "Public Synthetic Logistics", "2010 - 2024", "Coordinated shift schedules and trained colleagues", "Inventory reconciliation and stakeholder communication"),
}


def document_for(background="audit", *, source_id=None):
    source_id = source_id or str(uuid4())
    text = "\n".join(BACKGROUNDS[background])
    return ParsedResumeDocument(source_id, "docx", (
        ParsedResumeBlock(f"{source_id}:block:1", source_id, "docx", "body/1", "paragraph", text),
    ))


def evidence_output(document, *, category="work_experience", claim=None):
    block = document.blocks[0]
    lines = block.text.splitlines()
    item = {"evidence_id": "resume_evidence_001", "category": category,
        "normalized_claim": claim or lines[0], "source_block_ids": [block.block_id],
        "source_quotes": [{"block_id": block.block_id, "excerpt": block.text[:800]}],
        "confidence": "explicit", "uncertainty": "none", "evidence_origin": "resume_provided",
        "claim_type": "reported_fact"}
    if category == "work_experience":
        item.update(role_title=lines[0], organization=lines[1], time_range=lines[2],
                    responsibilities=[lines[3]], achievements=[], domain_signals=[], tools=[], business_metrics=[])
    if category == "projects":
        item.update(project_name=lines[0], responsibilities=[], achievements=[])
    return {"items": [item], "uncertainties": []}


def docx_for(background):
    document = Document()
    document.add_paragraph("\n".join(BACKGROUNDS[background]))
    output = BytesIO()
    document.save(output)
    return output.getvalue()
